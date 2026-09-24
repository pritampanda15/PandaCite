"""Offline tests: extractor HTTP handling, manager, BibTeX export and Word processing."""
from unittest import mock

import pytest
from docx import Document

from pandacite import EnhancedCitationManager
from pandacite.extractors import metadata as md_module
from pandacite.extractors.detector import IDDetector
from pandacite.processors.numbered import NumberedCitationProcessor
from pandacite.processors.word import CommandLineWordProcessor

CROSSREF = {"message": {
    "title": ["A title"], "author": [{"family": "Smith", "given": "John"}],
    "container-title": ["Nature"], "issued": {"date-parts": [[2020, 1]]},
    "volume": "5", "page": "1-10", "DOI": "10.1038/x",
}}


def fake_response(json_data=None, content=b""):
    response = mock.Mock()
    response.json.return_value = json_data
    response.content = content
    response.raise_for_status.return_value = None
    return response


@pytest.fixture
def manager():
    return EnhancedCitationManager()


def test_requests_have_timeout_and_user_agent():
    with mock.patch.object(md_module.requests, "get", return_value=fake_response(CROSSREF)) as get:
        md_module.EnhancedMetadataExtractor().extract_from_doi("10.1038/x")
    kwargs = get.call_args.kwargs
    assert kwargs["timeout"] == md_module.REQUEST_TIMEOUT
    assert "PandaCite" in kwargs["headers"]["User-Agent"]


def test_doi_url_is_normalized_and_year_falls_back_to_issued():
    with mock.patch.object(md_module.requests, "get", return_value=fake_response(CROSSREF)) as get:
        metadata = md_module.EnhancedMetadataExtractor().extract_from_doi("https://doi.org/10.1038/x")
    assert get.call_args.args[0] == "https://api.crossref.org/works/10.1038/x"
    assert metadata["year"] == "2020"
    assert metadata["authors"] == ["Smith, John"]


def test_crossref_empty_title_list():
    data = {"message": dict(CROSSREF["message"], title=[])}
    with mock.patch.object(md_module.requests, "get", return_value=fake_response(data)):
        metadata = md_module.EnhancedMetadataExtractor().extract_from_doi("10.1038/x")
    assert metadata["title"] == ""


def test_arxiv_prefix_is_not_mangled():
    with mock.patch.object(md_module.requests, "get", return_value=fake_response(content=b"<feed/>")) as get:
        md_module.EnhancedMetadataExtractor().extract_from_arxiv("arXiv:2101.00001v2")
    assert get.call_args.args[0].endswith("id_list=2101.00001")


def test_pubmed_url_with_trailing_slash():
    extractor = md_module.EnhancedMetadataExtractor()
    with mock.patch.object(extractor, "extract_from_pmid", return_value={}) as pmid:
        extractor.extract_from_url("https://pubmed.ncbi.nlm.nih.gov/12345678/")
    pmid.assert_called_once_with("12345678")


def test_bibtex_export_keys_are_valid_and_unique(manager, tmp_path):
    entry = {"authors": ["van der Berg, Ann"], "year": "2020", "title": "T"}
    manager.citation_data = {"a": dict(entry), "b": dict(entry)}
    out = tmp_path / "refs.bib"
    assert manager.export_bibtex(str(out))
    text = out.read_text(encoding="utf-8")
    assert "@article{vanderBerg2020," in text
    assert "@article{vanderBerg2020a," in text


def test_export_citations_writes_unicode(manager, tmp_path):
    out = tmp_path / "refs.txt"
    assert manager.export_citations(["Smith J, pp. 1–10"], str(out))
    assert "–" in out.read_text(encoding="utf-8")


def test_word_scan_ignores_bare_numbers(manager):
    processor = CommandLineWordProcessor(manager)
    citations = {}
    processor._process_text_for_citations(
        "In 2020 we saw [1] and (Smith, 2019). See PMID: 12345 and 10.1038/x.",
        citations, IDDetector())
    direct = {c["id_value"] for c in citations.values() if c["pattern"] == "direct_id"}
    assert direct == {"12345", "10.1038/x"}


def test_numbered_citations_follow_text_order(manager):
    doc = Document()
    doc.add_paragraph("First B then A.")
    citations = {"a": {"source_text": "A", "metadata_key": "ka"},
                 "b": {"source_text": "B", "metadata_key": "kb"}}
    order = NumberedCitationProcessor(manager).process_document(doc, citations, {}, "ieee")
    assert order == {"kb": 1, "ka": 2}


def test_word_document_end_to_end(manager, tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("As shown by 10.1038/x, pandas are great.")
    doc.save(src)

    processor = CommandLineWordProcessor(manager)
    document, citations = processor.process_document(str(src), "apa", IDDetector())
    metadata = {"doi-10.1038/x-metadata": {"authors": ["Smith, John"], "year": "2020",
                                           "title": "T", "journal": "J"}}
    for citation in citations.values():
        citation["metadata_key"] = "doi-10.1038/x-metadata"
    assert processor.update_document_with_citations(document, citations, metadata, "apa", str(out))
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "10.1038/x" not in text.split("References")[0]
    assert "References" in text
