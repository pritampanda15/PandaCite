"""Offline tests: CLI commands and exit codes."""
import sys
from unittest import mock

from docx import Document

from pandacite import cli

METADATA = {"authors": ["Smith, John"], "year": "2020", "title": "T", "journal": "J", "doi": "10.1038/x"}


def run_cli(*argv):
    with mock.patch.object(sys, "argv", ["pandacite", *argv]):
        return cli.main()


def test_single_success_and_bibtex_message(tmp_path, capsys):
    bib = tmp_path / "out.bib"
    with mock.patch("pandacite.extractors.metadata.EnhancedMetadataExtractor.extract_from_doi",
                    return_value=dict(METADATA)):
        assert run_cli("single", "-i", "10.1038/x", "-b", str(bib)) is None
    output = capsys.readouterr().out
    assert f"BibTeX exported to {bib}" in output
    assert "Failed" not in output
    assert "Smith2020" in bib.read_text()


def test_single_failure_exit_code():
    with mock.patch("pandacite.extractors.metadata.EnhancedMetadataExtractor.extract_from_doi",
                    return_value=None):
        assert run_cli("single", "-i", "10.1038/x") == 1


def test_batch_requires_input():
    assert run_cli("batch") == 2


def test_word_command_runs(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("See 10.1038/x for details.")
    doc.save(src)
    with mock.patch("pandacite.extractors.metadata.EnhancedMetadataExtractor.extract_from_doi",
                    return_value=dict(METADATA)):
        run_cli("word", "-i", str(src), "-o", str(out), "-f", "ieee")
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "See [1] for details." in text
    assert "1. Smith" in text


def test_word_command_merges_same_paper_cited_by_doi_and_pmid(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("By DOI (10.1038/x) and by PMID (PMID: 123).")
    doc.save(src)
    with mock.patch("pandacite.extractors.metadata.EnhancedMetadataExtractor.extract_from_doi",
                    return_value=dict(METADATA)), \
         mock.patch("pandacite.extractors.metadata.EnhancedMetadataExtractor.extract_from_pmid",
                    return_value=dict(METADATA)):
        run_cli("word", "-i", str(src), "-o", str(out), "-f", "apa")
    paragraphs = [p.text for p in Document(out).paragraphs if p.text]
    assert paragraphs[0] == "By DOI (Smith, 2020) and by PMID (Smith, 2020)."
    references = paragraphs[paragraphs.index("References") + 1:]
    assert references == ["Smith, J. (2020). T. J. https://doi.org/10.1038/x"]
